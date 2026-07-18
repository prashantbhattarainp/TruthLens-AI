#!/usr/bin/env python3
"""Read-only Phase 3.4 EDA for the immutable BharatFakeNewsKosh archive.

This program deliberately reads the ZIP-contained workbook in memory.  It never
extracts, edits, saves, normalises, filters, or otherwise creates a derivative
dataset.  It produces only aggregate metadata and vector research figures.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
import sys
import zipfile
from collections import Counter
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from statistics import fmean, stdev
from typing import Any, Iterable, Sequence
from xml.sax.saxutils import escape

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
RAW_ARCHIVE = ROOT / "ml" / "data" / "raw" / "bharatfakenewskosh-v1.zip"
METADATA_OUTPUT = ROOT / "ml" / "metadata" / "bharatfakenewskosh-v1.eda.json"
FIGURE_DIRECTORY = ROOT / "docs" / "research" / "figures"

DATASET_ID = "indian-digital-media-2026.07.17-r1"
REGISTRY_ID = "DSR-001"
PRIMARY_SHEET = "A"
ARTICLE_FIELD = "News Body"
HEADLINE_FIELD = "Statement"
RAW_TOKEN_PATTERN = re.compile(r"[^\W_]+(?:['’][^\W_]+)*", flags=re.UNICODE)
# This is a diagnostic signature, not a decoder: it detects common sequences
# produced when UTF-8 bytes are displayed as Latin-1-like characters.
POSSIBLE_MOJIBAKE_PATTERN = re.compile(r"(?:Ã.|Â.|à[\x80-\xBF])")
TEXT_FIELDS = [
    "Statement",
    "Eng_Trans_Statement",
    "News Body",
    "Eng_Trans_News_Body",
    "Text",
]

COLORS = {
    "blue": "#0072B2",
    "orange": "#D55E00",
    "green": "#009E73",
    "purple": "#6A3D9A",
    "neutral": "#6C757D",
    "dark": "#1F2937",
    "grid": "#D9E1E8",
    "pale_blue": "#D8ECF6",
    "pale_orange": "#FBE2D5",
    "background": "#FFFFFF",
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def is_literal_missing(value: Any) -> bool:
    """Treat only a blank cell or empty string as missing; do not trim values."""
    return value is None or (isinstance(value, str) and value == "")


def is_whitespace_only(value: Any) -> bool:
    return isinstance(value, str) and value != "" and value.strip() == ""


def raw_display(value: Any) -> str:
    """Display a stored value without normalising its spelling or case."""
    if value is None:
        return "<blank>"
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def raw_key(value: Any) -> tuple[str, str]:
    """Create a deterministic comparison key without changing the stored value."""
    if value is None:
        return ("none", "")
    if isinstance(value, datetime):
        return ("datetime", value.isoformat())
    if isinstance(value, date):
        return ("date", value.isoformat())
    return (type(value).__name__, str(value))


def raw_text(value: Any) -> str:
    return "" if value is None else str(value)


def raw_tokens(value: Any) -> list[str]:
    """Return raw-form word-like tokens; case and token forms are preserved."""
    return RAW_TOKEN_PATTERN.findall(raw_text(value))


def quantile(sorted_values: Sequence[float], probability: float) -> float:
    if not sorted_values:
        return 0.0
    position = (len(sorted_values) - 1) * probability
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return float(sorted_values[lower])
    fraction = position - lower
    return float(sorted_values[lower] * (1 - fraction) + sorted_values[upper] * fraction)


def number(value: float) -> int | float:
    rounded = round(float(value), 2)
    return int(rounded) if rounded.is_integer() else rounded


def describe(values: Sequence[int | float]) -> dict[str, int | float]:
    ordered = sorted(float(value) for value in values)
    if not ordered:
        return {"count": 0}
    return {
        "count": len(ordered),
        "min": number(ordered[0]),
        "q1": number(quantile(ordered, 0.25)),
        "median": number(quantile(ordered, 0.5)),
        "mean": number(fmean(ordered)),
        "q3": number(quantile(ordered, 0.75)),
        "p95": number(quantile(ordered, 0.95)),
        "p99": number(quantile(ordered, 0.99)),
        "max": number(ordered[-1]),
        "std_dev": number(stdev(ordered)) if len(ordered) > 1 else 0,
    }


def iqr_outliers(values: Sequence[int | float]) -> dict[str, int | float]:
    ordered = sorted(float(value) for value in values)
    if not ordered:
        return {"count": 0}
    first_quartile = quantile(ordered, 0.25)
    third_quartile = quantile(ordered, 0.75)
    iqr = third_quartile - first_quartile
    lower_fence = first_quartile - 1.5 * iqr
    upper_fence = third_quartile + 1.5 * iqr
    low_count = sum(value < lower_fence for value in ordered)
    high_count = sum(value > upper_fence for value in ordered)
    return {
        "count": len(ordered),
        "q1": number(first_quartile),
        "q3": number(third_quartile),
        "iqr": number(iqr),
        "lower_fence": number(lower_fence),
        "upper_fence": number(upper_fence),
        "low_outlier_count": low_count,
        "high_outlier_count": high_count,
        "total_outlier_count": low_count + high_count,
    }


def duplicate_metrics(keys: Iterable[Any]) -> dict[str, int]:
    counts = Counter(keys)
    duplicate_sizes = [count for count in counts.values() if count > 1]
    return {
        "distinct_values": len(counts),
        "duplicate_clusters": len(duplicate_sizes),
        "rows_in_duplicate_clusters": sum(duplicate_sizes),
        "extra_duplicate_rows": sum(count - 1 for count in duplicate_sizes),
        "largest_duplicate_cluster": max(duplicate_sizes, default=0),
    }


def parse_publish_date(value: Any) -> date | None:
    """Parse only typed Excel dates and unambiguous ISO-like string values."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        return None
    candidate = value.strip()
    if not candidate:
        return None
    try:
        return datetime.fromisoformat(candidate.replace("Z", "+00:00")).date()
    except ValueError:
        pass
    for pattern in ("%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(candidate, pattern).date()
        except ValueError:
            continue
    return None


def publish_date_representation(value: Any) -> str:
    """Classify a stored date representation without rewriting or parsing it."""
    if isinstance(value, datetime):
        return "typed_datetime"
    if isinstance(value, date):
        return "typed_date"
    if value is None:
        return "blank"
    if isinstance(value, str):
        if value == "":
            return "empty_string"
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:[T ].*)?", value.strip()):
            return "iso_like_string"
        if re.fullmatch(r"\d{4}[/.]\d{1,2}[/.]\d{1,2}", value.strip()):
            return "year_first_numeric_string"
        if re.fullmatch(r"\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4}", value.strip()):
            return "ambiguous_numeric_date_string"
        return "other_string"
    return type(value).__name__


def top_items(counter: Counter[str], limit: int | None = None) -> list[dict[str, int | str]]:
    items = sorted(counter.items(), key=lambda item: (-item[1], item[0]))
    if limit is not None:
        items = items[:limit]
    return [{"value": value, "count": count} for value, count in items]


def short_label(value: str, maximum: int = 48) -> str:
    return value if len(value) <= maximum else f"{value[: maximum - 1]}…"


def svg_document(title: str, description: str, width: int, height: int, elements: list[str]) -> str:
    style = f"""
<style>
  .title {{ font: 600 28px Arial, Helvetica, sans-serif; fill: {COLORS['dark']}; }}
  .subtitle {{ font: 15px Arial, Helvetica, sans-serif; fill: {COLORS['neutral']}; }}
  .axis {{ font: 13px Arial, Helvetica, sans-serif; fill: {COLORS['dark']}; }}
  .small {{ font: 12px Arial, Helvetica, sans-serif; fill: {COLORS['neutral']}; }}
  .value {{ font: 13px Arial, Helvetica, sans-serif; fill: {COLORS['dark']}; }}
  .axis-line {{ stroke: {COLORS['neutral']}; stroke-width: 1.2; }}
  .grid {{ stroke: {COLORS['grid']}; stroke-width: 1; }}
</style>"""
    return "\n".join(
        [
            f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}" role="img" aria-labelledby="chart-title chart-desc">',
            f"<title id=\"chart-title\">{escape(title)}</title>",
            f"<desc id=\"chart-desc\">{escape(description)}</desc>",
            f'<rect width="{width}" height="{height}" fill="{COLORS["background"]}"/>',
            style,
            *elements,
            "</svg>",
        ]
    )


def svg_text(x: float, y: float, content: str, css_class: str, anchor: str = "start") -> str:
    return f'<text x="{x:.1f}" y="{y:.1f}" class="{css_class}" text-anchor="{anchor}">{escape(content)}</text>'


def svg_rect(x: float, y: float, width: float, height: float, fill: str, opacity: float = 1.0) -> str:
    return f'<rect x="{x:.1f}" y="{y:.1f}" width="{max(width, 0):.1f}" height="{max(height, 0):.1f}" fill="{fill}" opacity="{opacity:.3f}"/>'


def write_figure(filename: str, title: str, description: str, width: int, height: int, elements: list[str]) -> None:
    FIGURE_DIRECTORY.mkdir(parents=True, exist_ok=True)
    (FIGURE_DIRECTORY / filename).write_text(
        svg_document(title, description, width, height, elements), encoding="utf-8"
    )


def horizontal_bar_chart(
    filename: str,
    title: str,
    subtitle: str,
    items: Sequence[tuple[str, int | float]],
    *,
    color: str = COLORS["blue"],
    unit: str = "records",
    zero_floor: bool = False,
) -> None:
    width = 1440
    height = max(560, 190 + len(items) * 38)
    left, right, top, bottom = 410, 170, 145, 80
    chart_width = width - left - right
    chart_height = height - top - bottom
    maximum = max((float(value) for _, value in items), default=0.0)
    scale_maximum = max(maximum, 1.0) if zero_floor else max(maximum, 1.0)
    elements = [
        svg_text(60, 54, title, "title"),
        svg_text(60, 82, subtitle, "subtitle"),
    ]
    for tick in range(6):
        value = scale_maximum * tick / 5
        x = left + chart_width * tick / 5
        elements.append(f'<line x1="{x:.1f}" y1="{top}" x2="{x:.1f}" y2="{top + chart_height}" class="grid"/>')
        elements.append(svg_text(x, height - 42, f"{value:,.0f}", "small", "middle"))
    elements.append(f'<line x1="{left}" y1="{top + chart_height}" x2="{left + chart_width}" y2="{top + chart_height}" class="axis-line"/>')
    row_height = chart_height / max(len(items), 1)
    for index, (label, raw_value) in enumerate(items):
        value = float(raw_value)
        y = top + index * row_height + row_height * 0.17
        bar_height = row_height * 0.64
        bar_width = chart_width * value / scale_maximum
        elements.append(svg_text(left - 16, y + bar_height * 0.72, short_label(str(label)), "axis", "end"))
        if value > 0:
            elements.append(svg_rect(left, y, bar_width, bar_height, color))
        else:
            elements.append(f'<circle cx="{left + 4:.1f}" cy="{y + bar_height / 2:.1f}" r="3.5" fill="{color}"/>')
        elements.append(svg_text(min(left + bar_width + 12, width - 26), y + bar_height * 0.72, f"{value:,.0f}", "value"))
    elements.append(svg_text(left + chart_width / 2, height - 16, unit, "axis", "middle"))
    write_figure(filename, title, f"{subtitle}. Horizontal bars report {unit}.", width, height, elements)


def vertical_bar_chart(
    filename: str,
    title: str,
    subtitle: str,
    items: Sequence[tuple[str, int | float]],
    *,
    colors: Sequence[str] | None = None,
    unit: str = "records",
) -> None:
    width, height = 1440, 820
    left, right, top, bottom = 130, 70, 145, 180
    chart_width, chart_height = width - left - right, height - top - bottom
    maximum = max((float(value) for _, value in items), default=1.0)
    maximum = max(maximum, 1.0)
    bar_slot = chart_width / max(len(items), 1)
    bar_width = bar_slot * 0.64
    palette = colors or [COLORS["blue"]]
    elements = [svg_text(60, 54, title, "title"), svg_text(60, 82, subtitle, "subtitle")]
    for tick in range(6):
        value = maximum * tick / 5
        y = top + chart_height - chart_height * tick / 5
        elements.append(f'<line x1="{left}" y1="{y:.1f}" x2="{left + chart_width}" y2="{y:.1f}" class="grid"/>')
        elements.append(svg_text(left - 12, y + 4, f"{value:,.0f}", "small", "end"))
    elements.append(f'<line x1="{left}" y1="{top}" x2="{left}" y2="{top + chart_height}" class="axis-line"/>')
    elements.append(f'<line x1="{left}" y1="{top + chart_height}" x2="{left + chart_width}" y2="{top + chart_height}" class="axis-line"/>')
    for index, (label, raw_value) in enumerate(items):
        value = float(raw_value)
        x = left + index * bar_slot + (bar_slot - bar_width) / 2
        bar_height = chart_height * value / maximum
        y = top + chart_height - bar_height
        color = palette[index % len(palette)]
        elements.append(svg_rect(x, y, bar_width, bar_height, color))
        elements.append(svg_text(x + bar_width / 2, y - 10, f"{value:,.0f}", "value", "middle"))
        label_y = top + chart_height + 30
        elements.append(f'<text x="{x + bar_width / 2:.1f}" y="{label_y:.1f}" class="axis" text-anchor="end" transform="rotate(-38 {x + bar_width / 2:.1f} {label_y:.1f})">{escape(short_label(str(label), 34))}</text>')
    elements.append(svg_text(left + chart_width / 2, height - 24, unit, "axis", "middle"))
    write_figure(filename, title, f"{subtitle}. Vertical bars report {unit}.", width, height, elements)


def histogram_chart(
    filename: str,
    title: str,
    subtitle: str,
    values: Sequence[int | float],
    *,
    x_label: str,
    color: str,
    log_scale: bool,
) -> None:
    width, height = 1440, 820
    left, right, top, bottom = 130, 60, 145, 100
    chart_width, chart_height = width - left - right, height - top - bottom
    filtered = [float(value) for value in values if float(value) >= 0]
    positive = [value for value in filtered if value > 0]
    if not positive:
        raise RuntimeError(f"Cannot draw histogram for empty values: {filename}")
    lower, upper = min(positive), max(positive)
    if lower == upper:
        upper = lower + 1
    bin_count = 48
    if log_scale:
        log_lower, log_upper = math.log10(lower), math.log10(upper)
        edges = [10 ** (log_lower + (log_upper - log_lower) * index / bin_count) for index in range(bin_count + 1)]
        project = lambda value: (math.log10(max(value, lower)) - log_lower) / (log_upper - log_lower)
        tick_values = [10 ** (log_lower + (log_upper - log_lower) * index / 4) for index in range(5)]
    else:
        edges = [lower + (upper - lower) * index / bin_count for index in range(bin_count + 1)]
        project = lambda value: (value - lower) / (upper - lower)
        tick_values = [lower + (upper - lower) * index / 4 for index in range(5)]
    counts = [0] * bin_count
    for value in positive:
        index = min(bin_count - 1, max(0, next((index for index in range(bin_count) if value < edges[index + 1]), bin_count - 1)))
        counts[index] += 1
    maximum = max(counts, default=1)
    elements = [svg_text(60, 54, title, "title"), svg_text(60, 82, subtitle, "subtitle")]
    for tick in range(6):
        value = maximum * tick / 5
        y = top + chart_height - chart_height * tick / 5
        elements.append(f'<line x1="{left}" y1="{y:.1f}" x2="{left + chart_width}" y2="{y:.1f}" class="grid"/>')
        elements.append(svg_text(left - 12, y + 4, f"{value:,.0f}", "small", "end"))
    for tick_value in tick_values:
        x = left + chart_width * project(tick_value)
        elements.append(f'<line x1="{x:.1f}" y1="{top + chart_height}" x2="{x:.1f}" y2="{top + chart_height + 6}" class="axis-line"/>')
        elements.append(svg_text(x, height - 52, f"{tick_value:,.0f}", "small", "middle"))
    elements.append(f'<line x1="{left}" y1="{top}" x2="{left}" y2="{top + chart_height}" class="axis-line"/>')
    elements.append(f'<line x1="{left}" y1="{top + chart_height}" x2="{left + chart_width}" y2="{top + chart_height}" class="axis-line"/>')
    for index, count in enumerate(counts):
        x_start = left + chart_width * project(edges[index])
        x_end = left + chart_width * project(edges[index + 1])
        height_value = chart_height * count / maximum
        elements.append(svg_rect(x_start + 0.8, top + chart_height - height_value, max(1, x_end - x_start - 1.6), height_value, color, 0.86))
    scale_label = " (logarithmic scale)" if log_scale else ""
    elements.append(svg_text(left + chart_width / 2, height - 16, f"{x_label}{scale_label}", "axis", "middle"))
    write_figure(filename, title, f"{subtitle}. Histogram uses {len(positive):,} non-empty raw values.", width, height, elements)


def boxplot_chart(
    filename: str,
    title: str,
    subtitle: str,
    series: Sequence[tuple[str, Sequence[int | float], str]],
    *,
    x_label: str,
) -> None:
    width, height = 1440, 620
    left, right, top, bottom = 270, 90, 135, 90
    chart_width, chart_height = width - left - right, height - top - bottom
    all_values = [float(value) for _, values, _ in series for value in values if float(value) > 0]
    minimum, maximum = min(all_values), max(all_values)
    log_minimum, log_maximum = math.log10(minimum), math.log10(maximum)
    if log_minimum == log_maximum:
        log_maximum += 1

    def x_position(value: float) -> float:
        return left + chart_width * (math.log10(max(value, minimum)) - log_minimum) / (log_maximum - log_minimum)

    elements = [svg_text(60, 54, title, "title"), svg_text(60, 82, subtitle, "subtitle")]
    for tick in range(5):
        value = 10 ** (log_minimum + (log_maximum - log_minimum) * tick / 4)
        x = x_position(value)
        elements.append(f'<line x1="{x:.1f}" y1="{top}" x2="{x:.1f}" y2="{top + chart_height}" class="grid"/>')
        elements.append(svg_text(x, height - 50, f"{value:,.0f}", "small", "middle"))
    elements.append(f'<line x1="{left}" y1="{top + chart_height}" x2="{left + chart_width}" y2="{top + chart_height}" class="axis-line"/>')
    row_height = chart_height / len(series)
    for index, (label, values, color) in enumerate(series):
        ordered = sorted(float(value) for value in values if float(value) > 0)
        row_y = top + row_height * (index + 0.5)
        minimum_value, maximum_value = ordered[0], ordered[-1]
        q1, median, q3 = quantile(ordered, 0.25), quantile(ordered, 0.5), quantile(ordered, 0.75)
        elements.append(svg_text(left - 18, row_y + 5, label, "axis", "end"))
        elements.append(f'<line x1="{x_position(minimum_value):.1f}" y1="{row_y:.1f}" x2="{x_position(maximum_value):.1f}" y2="{row_y:.1f}" stroke="{color}" stroke-width="3"/>')
        elements.append(f'<line x1="{x_position(minimum_value):.1f}" y1="{row_y - 14:.1f}" x2="{x_position(minimum_value):.1f}" y2="{row_y + 14:.1f}" stroke="{color}" stroke-width="2"/>')
        elements.append(f'<line x1="{x_position(maximum_value):.1f}" y1="{row_y - 14:.1f}" x2="{x_position(maximum_value):.1f}" y2="{row_y + 14:.1f}" stroke="{color}" stroke-width="2"/>')
        elements.append(svg_rect(x_position(q1), row_y - 26, x_position(q3) - x_position(q1), 52, color, 0.35))
        elements.append(f'<line x1="{x_position(median):.1f}" y1="{row_y - 30:.1f}" x2="{x_position(median):.1f}" y2="{row_y + 30:.1f}" stroke="{color}" stroke-width="4"/>')
        elements.append(svg_text(left + chart_width + 12, row_y + 5, f"median {median:,.0f}", "value"))
    elements.append(svg_text(left + chart_width / 2, height - 16, f"{x_label} (logarithmic scale)", "axis", "middle"))
    write_figure(filename, title, f"{subtitle}. Each box spans the first to third quartile; the heavy line marks the median.", width, height, elements)


def load_raw_workbook() -> tuple[list[dict[str, Any]], list[str], dict[str, Any], dict[str, int], str, str]:
    if not RAW_ARCHIVE.exists():
        raise FileNotFoundError(f"Raw archive is missing: {RAW_ARCHIVE}")
    archive_sha = sha256_file(RAW_ARCHIVE)
    with zipfile.ZipFile(RAW_ARCHIVE) as archive:
        members = archive.namelist()
        if len(members) != 1 or not members[0].lower().endswith(".xlsx"):
            raise RuntimeError("Expected exactly one XLSX member in the immutable raw archive")
        workbook_member = members[0]
        workbook_bytes = archive.read(workbook_member)
        workbook_sha = sha256_bytes(workbook_bytes)
    values_workbook = openpyxl.load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    formulas_workbook = openpyxl.load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=False)
    if PRIMARY_SHEET not in values_workbook.sheetnames:
        raise RuntimeError(f"Primary sheet {PRIMARY_SHEET!r} is not present")

    values_sheet = values_workbook[PRIMARY_SHEET]
    headers = [str(value) if value is not None else "" for value in next(values_sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    if len(headers) != len(set(headers)) or any(not header for header in headers):
        raise RuntimeError("Primary-sheet headers are missing or not unique")
    rows: list[dict[str, Any]] = []
    for values in values_sheet.iter_rows(min_row=2, values_only=True):
        if any(not is_literal_missing(value) for value in values):
            rows.append(dict(zip(headers, values, strict=True)))

    # OpenPyXL leaves an entirely empty read-only worksheet dimension undefined.
    # Count primary records from the already-read data and treat an undefined
    # auxiliary dimension as an empty worksheet rather than iterating Excel's
    # full row grid.
    sheet_composition: dict[str, Any] = {}
    for sheet_name in values_workbook.sheetnames:
        worksheet = values_workbook[sheet_name]
        if sheet_name == PRIMARY_SHEET:
            nonempty_rows = len(rows) + 1  # Declared header plus raw records.
        else:
            try:
                worksheet.calculate_dimension(force=True)
            except UnboundLocalError:
                nonempty_rows = 0
            else:
                nonempty_rows = sum(
                    any(not is_literal_missing(value) for value in row)
                    for row in worksheet.iter_rows(values_only=True)
                )
        sheet_composition[sheet_name] = {
            "nonempty_rows_including_header_if_present": nonempty_rows,
            "declared_role": "primary_data" if sheet_name == PRIMARY_SHEET else "auxiliary_or_empty",
        }

    formula_counts: Counter[str] = Counter()
    formula_sheet = formulas_workbook[PRIMARY_SHEET]
    for row in formula_sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.data_type == "f":
                formula_counts[headers[cell.column - 1]] += 1

    values_workbook.close()
    formulas_workbook.close()
    return rows, headers, sheet_composition, dict(sorted(formula_counts.items())), archive_sha, workbook_sha


def build_analysis() -> dict[str, Any]:
    rows, headers, sheet_composition, formula_counts, archive_sha_before, workbook_sha = load_raw_workbook()
    total_rows = len(rows)
    if total_rows != 26232:
        raise RuntimeError(f"Unexpected primary row count: {total_rows}")

    missingness: dict[str, dict[str, int | float]] = {}
    for header in headers:
        literal_missing = sum(is_literal_missing(row[header]) for row in rows)
        whitespace_only = sum(is_whitespace_only(row[header]) for row in rows)
        missingness[header] = {
            "literal_missing_count": literal_missing,
            "literal_missing_pct": number(100 * literal_missing / total_rows),
            "whitespace_only_nonempty_count": whitespace_only,
        }

    label_counter = Counter(raw_display(row["Label"]) for row in rows)
    recognised_raw_labels = {"True", "False"}
    invalid_raw_labels = sum(count for label, count in label_counter.items() if label not in recognised_raw_labels)
    label_percentages = {
        label: number(100 * count / total_rows) for label, count in sorted(label_counter.items())
    }

    article_values = [raw_text(row[ARTICLE_FIELD]) for row in rows]
    headline_values = [raw_text(row[HEADLINE_FIELD]) for row in rows]
    article_characters = [len(value) for value in article_values]
    headline_characters = [len(value) for value in headline_values]
    article_word_counts: list[int] = []
    headline_word_counts: list[int] = []
    article_vocabulary: set[str] = set()
    headline_vocabulary: set[str] = set()
    raw_token_counter: Counter[str] = Counter()
    for article, headline in zip(article_values, headline_values, strict=True):
        article_tokens = raw_tokens(article)
        headline_tokens = raw_tokens(headline)
        article_word_counts.append(len(article_tokens))
        headline_word_counts.append(len(headline_tokens))
        article_vocabulary.update(article_tokens)
        headline_vocabulary.update(headline_tokens)
        raw_token_counter.update(article_tokens)

    def counter_for(field: str) -> Counter[str]:
        return Counter(raw_display(row[field]) for row in rows if not is_literal_missing(row[field]))

    fact_check_source_counter = counter_for("Fact_Check_Source")
    source_type_counter = counter_for("Source_Type")
    language_counter = counter_for("Language")
    category_counter = counter_for("News_Category")
    platform_counter = counter_for("Platform")

    duplicate_results = {
        "id": duplicate_metrics(raw_key(row["id"]) for row in rows),
        "complete_primary_row": duplicate_metrics(
            tuple(raw_key(row[header]) for header in headers) for row in rows
        ),
        "statement": duplicate_metrics(raw_key(row[HEADLINE_FIELD]) for row in rows),
        "news_body": duplicate_metrics(raw_key(row[ARTICLE_FIELD]) for row in rows),
        "statement_news_body_pair": duplicate_metrics(
            (raw_key(row[HEADLINE_FIELD]), raw_key(row[ARTICLE_FIELD])) for row in rows
        ),
    }

    raw_dates = [row["Publish_Date"] for row in rows]
    parsed_dates = [parse_publish_date(value) for value in raw_dates]
    valid_dates = [value for value in parsed_dates if value is not None]
    date_counter = Counter(str(value.year) for value in valid_dates)
    date_missing_count = sum(is_literal_missing(value) for value in raw_dates)
    date_nonmissing_unparsed_count = sum(
        not is_literal_missing(raw) and parsed is None for raw, parsed in zip(raw_dates, parsed_dates, strict=True)
    )
    year_counts = top_items(date_counter)
    source_total = sum(fact_check_source_counter.values())
    source_shares = [count / source_total for count in fact_check_source_counter.values()] if source_total else []
    hhi = sum(share * share for share in source_shares)
    largest_source = top_items(fact_check_source_counter, 1)
    date_representation_counter = Counter(publish_date_representation(value) for value in raw_dates)
    encoding_profile: dict[str, dict[str, int | float]] = {}
    for field in TEXT_FIELDS:
        values = [raw_text(row[field]) for row in rows]
        possible_mojibake_records = sum(
            bool(POSSIBLE_MOJIBAKE_PATTERN.search(value)) for value in values
        )
        non_ascii_records = sum(any(ord(character) > 127 for character in value) for value in values)
        encoding_profile[field] = {
            "possible_utf8_as_latin1_pattern_record_count": possible_mojibake_records,
            "possible_utf8_as_latin1_pattern_record_pct": number(100 * possible_mojibake_records / total_rows),
            "non_ascii_record_count": non_ascii_records,
            "non_ascii_record_pct": number(100 * non_ascii_records / total_rows),
        }

    analysis_time = datetime.now(timezone.utc).isoformat()
    figures = [
        "01_label_distribution.svg",
        "02_dataset_composition.svg",
        "03_article_length_histogram.svg",
        "04_headline_length_histogram.svg",
        "05_word_count_distribution.svg",
        "06_character_count_distribution.svg",
        "07_missing_values.svg",
        "08_duplicate_analysis.svg",
        "09_source_distribution.svg",
        "10_top_raw_word_tokens.svg",
        "11_language_distribution.svg",
        "12_category_distribution.svg",
        "13_publication_year_distribution.svg",
        "14_possible_encoding_artifacts.svg",
    ]

    metadata: dict[str, Any] = {
        "schema_version": "1.0.0",
        "record_type": "dataset_eda",
        "analysis_id": "EDA-2026-07-17-DSR-001-R1",
        "dataset_id": DATASET_ID,
        "registry_id": REGISTRY_ID,
        "analysed_at": analysis_time,
        "analysis_mode": "read_only_no_extraction_no_transformation_no_derivative_dataset",
        "analysis_scope": {
            "included_raw_artifact": "ml/data/raw/bharatfakenewskosh-v1.zip",
            "primary_analysis_unit": "Each nonempty row in workbook sheet A after its declared header row.",
            "primary_row_count": total_rows,
            "excluded_from_primary_statistics": "Sheet1 is retained as undocumented headerless auxiliary content and Sheet3 is empty; neither is deleted, transformed, or represented as an analytical record.",
            "formula_representation": "Formula locations were counted from formula cells. Descriptive field statistics use only cached displayed values supplied in the raw workbook; formulas were not evaluated, recalculated, or written back.",
            "text_measurement": "Word-like tokens use the case-preserving Unicode pattern [^\\W_]+(?:['’][^\\W_]+)*. No lowercasing, stopword removal, stemming, lemmatisation, translation, normalisation, filtering, or text rewrite was performed.",
        },
        "reproducibility": {
            "script": "scripts/research/run_phase_3_4_eda.py",
            "script_sha256": sha256_file(Path(__file__)),
            "python_version": sys.version.split()[0],
            "openpyxl_version": openpyxl.__version__,
            "raw_archive_sha256_before_analysis": archive_sha_before,
            "source_workbook_sha256": workbook_sha,
        },
        "workbook_profile": {
            "sheet_names": list(sheet_composition),
            "sheet_composition": sheet_composition,
            "primary_sheet": PRIMARY_SHEET,
            "primary_headers": headers,
            "formula_cell_count_by_column": formula_counts,
        },
        "sample_statistics": {
            "total_primary_samples": total_rows,
            "raw_label_counts": dict(sorted(label_counter.items())),
            "raw_label_percentages": label_percentages,
            "raw_label_values_recognised_by_syntax": invalid_raw_labels == 0,
            "invalid_or_unexpected_raw_label_count": invalid_raw_labels,
            "label_mapping_status": "pending; raw True/False values have not been mapped to canonical REAL/FAKE labels",
            "dataset_wise_distribution": [{"dataset_id": DATASET_ID, "count": total_rows}],
        },
        "text_profile": {
            "article_field": ARTICLE_FIELD,
            "headline_field": HEADLINE_FIELD,
            "article_character_count": describe(article_characters),
            "headline_character_count": describe(headline_characters),
            "article_word_count": describe(article_word_counts),
            "headline_word_count": describe(headline_word_counts),
            "article_word_count_iqr_outliers": iqr_outliers(article_word_counts),
            "headline_word_count_iqr_outliers": iqr_outliers(headline_word_counts),
            "article_vocabulary_size_case_preserving": len(article_vocabulary),
            "headline_vocabulary_size_case_preserving": len(headline_vocabulary),
            "combined_vocabulary_size_case_preserving": len(article_vocabulary | headline_vocabulary),
            "top_raw_article_tokens_case_preserving_no_stopword_removal": top_items(raw_token_counter, 25),
        },
        "encoding_profile": {
            "container_encoding_check": "The Phase 3.3 XLSX container check found XML and relationship members UTF-8-decodable.",
            "text_representation_check": "The diagnostic below flags possible UTF-8-as-Latin-1-like character sequences in stored text. It does not repair, decode, or alter them.",
            "field_results": encoding_profile,
        },
        "missingness": missingness,
        "duplicate_analysis": {
            "method": "Exact equality only, using raw stored values (and cached formula values where present). No normalisation, near-duplicate detection, clustering, or removal was performed.",
            "results": duplicate_results,
        },
        "categorical_distributions": {
            "fact_check_source": top_items(fact_check_source_counter),
            "source_type": top_items(source_type_counter),
            "language": top_items(language_counter),
            "news_category": top_items(category_counter),
            "platform": top_items(platform_counter),
            "source_concentration": {
                "nonmissing_source_records": source_total,
                "distinct_fact_check_sources": len(fact_check_source_counter),
                "largest_source": largest_source[0] if largest_source else None,
                "largest_source_share_pct": number(100 * largest_source[0]["count"] / source_total) if largest_source else 0,
                "herfindahl_hirschman_index": number(hhi),
            },
        },
        "temporal_profile": {
            "publish_date_field": "Publish_Date",
            "typed_or_unambiguous_iso_date_count": len(valid_dates),
            "literal_missing_date_count": date_missing_count,
            "nonmissing_unparsed_date_count": date_nonmissing_unparsed_count,
            "stored_value_representation_counts": dict(sorted(date_representation_counter.items())),
            "earliest_parsed_date": min(valid_dates).isoformat() if valid_dates else None,
            "latest_parsed_date": max(valid_dates).isoformat() if valid_dates else None,
            "year_distribution": year_counts,
        },
        "figure_files": [f"docs/research/figures/{filename}" for filename in figures],
        "status": "descriptive_raw_eda_complete; semantic label, formula provenance, auxiliary-sheet purpose, and any derivative/split decisions remain pending",
    }

    # Publication-ready, editable SVG figures use only aggregate values.
    vertical_bar_chart(
        "01_label_distribution.svg",
        "Raw label distribution in primary sheet A",
        "Unmapped source labels; binary semantic mapping remains pending (n = 26,232)",
        [(label, count) for label, count in sorted(label_counter.items())],
        colors=[COLORS["blue"], COLORS["orange"]],
    )
    vertical_bar_chart(
        "02_dataset_composition.svg",
        "Unmodified workbook composition",
        "Nonempty rows by raw worksheet; sheet A includes its header row in this workbook-level view",
        [(name, profile["nonempty_rows_including_header_if_present"]) for name, profile in sheet_composition.items()],
        colors=[COLORS["blue"], COLORS["orange"], COLORS["neutral"]],
    )
    histogram_chart(
        "03_article_length_histogram.svg",
        "Article-field length distribution",
        "Raw case-preserving token counts in News Body (n = 26,232)",
        article_word_counts,
        x_label="Raw News Body token count",
        color=COLORS["blue"],
        log_scale=True,
    )
    histogram_chart(
        "04_headline_length_histogram.svg",
        "Headline-field length distribution",
        "Raw case-preserving token counts in Statement (n = 26,232)",
        headline_word_counts,
        x_label="Raw Statement token count",
        color=COLORS["orange"],
        log_scale=True,
    )
    boxplot_chart(
        "05_word_count_distribution.svg",
        "Raw word-count distribution",
        "No text preprocessing; boxes show quartiles and bold marks show medians",
        [
            ("News Body", article_word_counts, COLORS["blue"]),
            ("Statement", headline_word_counts, COLORS["orange"]),
        ],
        x_label="Raw case-preserving token count",
    )
    boxplot_chart(
        "06_character_count_distribution.svg",
        "Raw character-count distribution",
        "Stored character counts, including original whitespace and punctuation",
        [
            ("News Body", article_characters, COLORS["blue"]),
            ("Statement", headline_characters, COLORS["orange"]),
        ],
        x_label="Stored character count",
    )
    horizontal_bar_chart(
        "07_missing_values.svg",
        "Literal missing cells by primary-sheet field",
        "Blank cells or exact empty strings; whitespace-only nonempty values are reported separately in metadata",
        [(header, int(missingness[header]["literal_missing_count"])) for header in headers],
        color=COLORS["purple"],
        unit="literal missing cells",
        zero_floor=True,
    )
    duplicate_bars = [
        ("Duplicate IDs (extra rows)", duplicate_results["id"]["extra_duplicate_rows"]),
        ("Exact duplicate complete rows (extra rows)", duplicate_results["complete_primary_row"]["extra_duplicate_rows"]),
        ("Duplicate Statements (extra rows)", duplicate_results["statement"]["extra_duplicate_rows"]),
        ("Duplicate News Body values (extra rows)", duplicate_results["news_body"]["extra_duplicate_rows"]),
        ("Duplicate Statement + News Body pairs (extra rows)", duplicate_results["statement_news_body_pair"]["extra_duplicate_rows"]),
    ]
    horizontal_bar_chart(
        "08_duplicate_analysis.svg",
        "Exact duplicate analysis",
        "Raw equality only; counts are retained for review and no row was removed",
        duplicate_bars,
        color=COLORS["orange"],
        unit="extra duplicate rows",
        zero_floor=True,
    )
    horizontal_bar_chart(
        "09_source_distribution.svg",
        "Fact-check source distribution",
        "Top raw Fact_Check_Source values in sheet A",
        [(item["value"], int(item["count"])) for item in top_items(fact_check_source_counter, 15)],
        color=COLORS["green"],
        unit="records",
    )
    horizontal_bar_chart(
        "10_top_raw_word_tokens.svg",
        "Most frequent raw News Body tokens",
        "Case-preserving word-like tokens; no lowercasing or stopword removal",
        [(item["value"], int(item["count"])) for item in top_items(raw_token_counter, 20)],
        color=COLORS["blue"],
        unit="token occurrences",
    )
    horizontal_bar_chart(
        "11_language_distribution.svg",
        "Source-provided language distribution",
        "Raw Language values in sheet A",
        [(item["value"], int(item["count"])) for item in top_items(language_counter)],
        color=COLORS["purple"],
        unit="records",
    )
    horizontal_bar_chart(
        "12_category_distribution.svg",
        "News category distribution",
        "Top raw News_Category values in sheet A; includes cached values of source formula cells",
        [(item["value"], int(item["count"])) for item in top_items(category_counter, 15)],
        color=COLORS["orange"],
        unit="records",
    )
    vertical_bar_chart(
        "13_publication_year_distribution.svg",
        "Publish-date year distribution",
        "Typed Excel or unambiguous ISO-like dates only; unparsed nonblank dates are not silently assigned",
        [(item["value"], int(item["count"])) for item in year_counts],
        colors=[COLORS["green"]],
    )
    horizontal_bar_chart(
        "14_possible_encoding_artifacts.svg",
        "Possible text-encoding artifact indicators",
        "Records containing a UTF-8-as-Latin-1-like character signature; diagnostic only, no text was decoded or changed",
        [
            (field, int(profile["possible_utf8_as_latin1_pattern_record_count"]))
            for field, profile in encoding_profile.items()
        ],
        color=COLORS["purple"],
        unit="records with diagnostic signature",
        zero_floor=True,
    )

    archive_sha_after = sha256_file(RAW_ARCHIVE)
    if archive_sha_before != archive_sha_after:
        raise RuntimeError("Raw archive checksum changed during a read-only analysis; refusing to publish output")
    metadata["reproducibility"]["raw_archive_sha256_after_analysis"] = archive_sha_after
    METADATA_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    METADATA_OUTPUT.write_text(json.dumps(metadata, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return metadata


if __name__ == "__main__":
    result = build_analysis()
    print(
        json.dumps(
            {
                "analysis_id": result["analysis_id"],
                "primary_samples": result["sample_statistics"]["total_primary_samples"],
                "figures": len(result["figure_files"]),
                "metadata": str(METADATA_OUTPUT.relative_to(ROOT)),
            },
            indent=2,
        )
    )
